"""Deck text: python-pptx (MIT) for PPTX; pdftotext -layout for PDF (SPEC §6.6)."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path
from typing import Any, Optional


def _words(text: str) -> list[str]:
    return [tok for tok in text.replace("\n", " ").split() if tok]


def _bullets(lines: list[str]) -> int:
    count = 0
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped[:1] in {"•", "·", "-", "*", "–"}:
            count += 1
        elif len(stripped) < 80 and stripped[:2].rstrip(".").isdigit():
            count += 1
    return count


def extract_pptx(path: Path) -> list[dict[str, Any]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        bodies: list[str] = []
        notes = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape == slide.shapes.title:
                continue
            text = shape.text_frame.text.strip() if shape.text_frame else ""
            if text:
                bodies.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(bodies)
        lines = [ln for ln in body.splitlines() if ln.strip()]
        slides.append(
            {
                "index": i,
                "title": title,
                "body": body,
                "notes": notes,
                "text": "\n".join(p for p in (title, body) if p),
                "word_count": len(_words(title + " " + body)),
                "bullet_count": _bullets(lines),
                "source": "python-pptx",
            }
        )
    return slides


def extract_pdf(path: Path) -> list[dict[str, Any]]:
    exe = shutil.which("pdftotext")
    if not exe:
        raise FileNotFoundError("pdftotext is not installed")
    result = subprocess.run(
        [exe, "-layout", "-enc", "UTF-8", str(path), "-"],
        capture_output=True,
        timeout=60,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.decode("utf-8", "replace")[:200])
    raw = result.stdout.decode("utf-8", "replace")
    pages = raw.split("\f")
    slides: list[dict[str, Any]] = []
    for i, page in enumerate(pages, start=1):
        text = page.strip()
        if not text and i == len(pages):
            continue
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        title = lines[0] if lines else ""
        body = "\n".join(lines[1:])
        slides.append(
            {
                "index": i,
                "title": title,
                "body": body,
                "notes": "",
                "text": text,
                "word_count": len(_words(text)),
                "bullet_count": _bullets(lines),
                "source": "pdftotext",
            }
        )
    return slides


def extract_deck(path: Optional[Path]) -> Optional[dict[str, Any]]:
    if path is None or not path.exists():
        return None
    suffix = path.suffix.casefold()
    if suffix == ".pptx":
        pages = extract_pptx(path)
        kind = "pptx"
    elif suffix == ".pdf":
        pages = extract_pdf(path)
        kind = "pdf"
    else:
        return None
    return {
        "kind": kind,
        "path": str(path),
        "page_count": len(pages),
        "pages": pages,
    }
