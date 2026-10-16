"""Plan deck review (issue #60). No live slides_v1."""

from pathlib import Path

from pipeline.prep_slides import (
    ANY_WORD_FLAG,
    flag_slide,
    review_deck,
    review_slides,
    suggest_assertion_title,
)


def test_table_has_slide_issue_fix_and_assertion_title():
    slides = [
        {
            "index": 1,
            "title": "Agenda",
            "body": "Welcome\nTeam\nGoals",
            "text": "Agenda Welcome Team Goals",
            "word_count": 4,
            "bullet_count": 3,
        },
        {
            "index": 2,
            "title": "Revenue grew 12 percent this quarter",
            "body": "One chart.",
            "text": "Revenue grew 12 percent this quarter One chart.",
            "word_count": 8,
            "bullet_count": 0,
        },
    ]
    body = review_slides(slides)
    assert body["prompt_id"] == "prep_slides_v1"
    assert [row["slide"] for row in body["table"]] == [1, 2]
    for row in body["table"]:
        assert {"slide", "issue", "fix", "suggested_assertion_title"} <= set(row)
    assert "topic title" in body["table"][0]["issue"]
    assert body["table"][1]["issue"] == "ok"
    assert body["table"][0]["suggested_assertion_title"].endswith(".")


def test_storyline_and_minimal_text_heuristics():
    dense = {
        "index": 1,
        "title": "Overview",
        "body": " ".join(["word"] * 45),
        "text": "Overview " + " ".join(["word"] * 45),
        "word_count": 46,
        "bullet_count": 5,
    }
    flags = flag_slide(dense)
    assert any("25" in item or "40" in item for item in flags)
    assert any("bullet" in item for item in flags)
    assert ANY_WORD_FLAG == 40
    story = review_slides(
        [
            {"index": 1, "title": "The status quo is failing customers", "body": "", "word_count": 6, "bullet_count": 0},
            {"index": 2, "title": "Approve the pilot this quarter", "body": "", "word_count": 5, "bullet_count": 0},
        ]
    )["storyline"]
    assert story["passes"] is True
    assert "→" in story["text"]


def test_review_deck_reads_pptx(tmp_path: Path):
    from pptx import Presentation

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Q3"
    slide.placeholders[1].text = "More than twenty five filler words " + ("text " * 20)
    second = prs.slides.add_slide(layout)
    second.shapes.title.text = "We shipped the feature this quarter"
    second.placeholders[1].text = "One chart"
    prs.save(str(path))
    body = review_deck(path)
    assert len(body["table"]) == 2
    assert body["table"][0]["suggested_assertion_title"]
    assert suggest_assertion_title({"title": "Q3", "body": "x"}).endswith(".")
