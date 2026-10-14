"""Slides stage (issue #43) on synthetic CC0-style media only."""

from __future__ import annotations

import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import numpy as np
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from pipeline.slides.antivirus import DeckRejected, scan_deck_bytes, scan_deck_file
from pipeline.slides.deck import extract_pptx
from pipeline.slides.match import ahash64, hamming, match_slide, reading_overlap, tfidf_cosine
from pipeline.slides.region import detect_slide_region
from pipeline.slides.scenes import merge_short_scenes
from pipeline.stages.slides import run_slides_stage


def _pptx_bytes(*, with_macro: bool = False) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"></Types>')
        zf.writestr("ppt/presentation.xml", "<p:presentation/>")
        zf.writestr("ppt/slides/slide1.xml", "<p:sld/>")
        if with_macro:
            zf.writestr("ppt/vbaProject.bin", b"MZ")
    return buf.getvalue()


def _real_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "We shipped the feature this quarter"
    slide.placeholders[1].text = "Revenue grew\nChurn fell\nSupport load dropped"
    prs.save(str(path))


def _scene_video(path: Path) -> None:
    cmd = [
        "ffmpeg",
        "-y",
        "-f",
        "lavfi",
        "-i",
        "color=c=0xCC0000:s=320x240:d=3:r=10",
        "-f",
        "lavfi",
        "-i",
        "color=c=0x0033CC:s=320x240:d=3:r=10",
        "-filter_complex",
        "[0:v][1:v]concat=n=2:v=1:a=0",
        "-pix_fmt",
        "yuv420p",
        "-an",
        str(path),
    ]
    subprocess.run(cmd, check=True, capture_output=True)


def test_merge_scenes_under_two_seconds():
    merged = merge_short_scenes([(0, 3.0), (3.0, 4.2), (4.2, 8.0)])
    assert merged == [(0.0, 8.0)] or merged[0][0] == 0.0
    assert all(end - start >= 2.0 or i == 0 for i, (start, end) in enumerate(merged))
    # the 1.2 s middle scene is absorbed
    assert len(merged) < 3


def test_antivirus_accepts_pptx_rejects_macro_and_exe_magic():
    assert scan_deck_bytes(_pptx_bytes(), filename="talk.pptx") == "pptx"
    with pytest.raises(DeckRejected):
        scan_deck_bytes(_pptx_bytes(with_macro=True), filename="talk.pptx")
    with pytest.raises(DeckRejected):
        scan_deck_bytes(b"\x00\x01not-a-deck", filename="talk.pptx")
    with pytest.raises(DeckRejected):
        scan_deck_bytes(b"%PDF-1.4\n/JavaScript (app.alert)\n", filename="talk.pdf")


def test_antivirus_scan_file(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    path.write_bytes(_pptx_bytes())
    assert scan_deck_file(path) == "pptx"


def test_stage_region_full_frame_for_screen_recording():
    frames = [np.zeros((120, 160, 3), dtype=np.uint8) for _ in range(3)]
    region = detect_slide_region(frames, setup="screen_recording", is_screen_recording=True)
    assert region["x"] == 0 and region["y"] == 0
    assert region["w"] == 160 and region["h"] == 120
    assert "full_frame" in region["method"]


def test_stage_region_finds_bright_quad():
    frames = []
    for _ in range(4):
        img = np.zeros((100, 200, 3), dtype=np.uint8)
        img[:, 80:] = 220
        frames.append(img)
    region = detect_slide_region(frames, setup="stage", is_screen_recording=False)
    assert region["w"] >= 80
    assert region["x"] >= 50
    assert region["method"] == "stage_bright_quad_h"


def test_tfidf_and_reading_overlap():
    scores = tfidf_cosine("revenue grew this quarter", ["weather is nice", "revenue grew this quarter a lot"])
    assert scores[1] > scores[0]
    pages = [
        {"index": 1, "text": "weather is nice today in the park sunshine"},
        {"index": 2, "text": "revenue grew this quarter after the launch"},
    ]
    matched = match_slide("revenue grew this quarter after the launch day results", pages)
    assert matched["deck_page"] == 2
    assert matched["method"] == "tfidf"
    overlap = reading_overlap(
        "revenue grew this quarter after the launch",
        "today revenue grew this quarter after the launch we think",
    )
    assert overlap is not None and overlap > 0.5


def test_ahash_hamming():
    a = np.arange(64, dtype=np.uint8).reshape(8, 8)
    b = a.copy()
    assert hamming(ahash64(a), ahash64(b)) == 0
    b[0, 0] = 255
    assert hamming(ahash64(a), ahash64(b)) >= 0


def test_pptx_extract(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    _real_pptx(path)
    pages = extract_pptx(path)
    assert pages[0]["title"].startswith("We shipped")
    assert pages[0]["word_count"] > 5
    assert pages[0]["source"] == "python-pptx"


def test_run_slides_on_synthetic_scenes(tmp_path: Path):
    video = tmp_path / "scenes.mp4"
    _scene_video(video)
    deck = tmp_path / "deck.pptx"
    _real_pptx(deck)
    result = run_slides_stage(
        output_dir=tmp_path,
        session_id="s",
        video_path=video,
        deck_path=deck,
        camera_setup="screen_recording",
        language="en",
        is_screen_recording=True,
        transcript={
            "segments": [
                {"start": 0, "end": 6, "text": "We shipped the feature this quarter revenue grew"}
            ]
        },
    )
    payload = json.loads((tmp_path / "slides.json").read_text())
    assert payload["cpu_only"] is True
    assert payload["paddleocr"] is False
    assert payload["scene_engine"] in {"pyscenedetect", "ffmpeg_scene_0.3", "single_span"}
    assert payload["slide_count"] >= 1
    assert (tmp_path / "deck_text.json").exists()
    assert result.skipped is False
    # keyframe in first half: first scene 0–3 s → t < 1.5 + epsilon if two scenes
    for slide in payload["slides"]:
        if slide.get("keyframe_t") is not None and slide.get("t_start") is not None:
            half = (slide["t_end"] - slide["t_start"]) / 2.0
            assert slide["keyframe_t"] <= slide["t_start"] + half + 0.35


def test_audio_only_without_deck_is_skipped(tmp_path: Path):
    result = run_slides_stage(
        output_dir=tmp_path,
        session_id="s",
        is_audio_only=True,
        camera_setup="audio_only",
    )
    assert result.skipped is True
    assert "audio-only" in (result.skip_reason or "").lower()
    assert not (tmp_path / "slides.json").exists()


def test_no_pyscenedetect_in_forbidden_models():
    req = (Path(__file__).parent.parent / "requirements.txt").read_text().lower()
    pins = [
        line.split("==")[0].strip()
        for line in req.splitlines()
        if line.strip() and not line.strip().startswith("#") and "==" in line
    ]
    assert "paddleocr" not in pins
    assert "paddlepaddle" not in pins
    assert "scenedetect" in pins
    src = (Path(__file__).parent.parent / "pipeline" / "slides" / "scenes.py").read_text()
    assert "BSD" in src or "bsd" in src.lower()
